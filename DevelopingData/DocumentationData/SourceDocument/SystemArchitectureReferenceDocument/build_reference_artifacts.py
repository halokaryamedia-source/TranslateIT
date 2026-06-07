from __future__ import annotations

import hashlib
import json
import os
import re
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


THIS_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = THIS_DIR.parents[3]
MARKDOWN_PATH = THIS_DIR / "TRANSLATEIT_SYSTEM_ARCHITECTURE_AND_ENGINE_REFERENCE.md"
DOCX_PATH = THIS_DIR / "TRANSLATEIT_SYSTEM_ARCHITECTURE_AND_ENGINE_REFERENCE.docx"
MANIFEST_PATH = THIS_DIR / "TRANSLATEIT_SYSTEM_ARCHITECTURE_MANIFEST.json"


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _file_record(path: Path, root: Path) -> dict[str, object]:
    stat = path.stat()
    return {
        "path": str(path.relative_to(root)).replace("\\", "/"),
        "absolute_path": str(path),
        "bytes": stat.st_size,
        "sha256": _sha256(path),
        "modified_utc": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(timespec="seconds"),
    }


def _set_cell_text(cell, text: str, *, bold: bool = False, size: int = 10, color: str = "000000") -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Arial"


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_width(cell, width_in: float) -> None:
    cell.width = Inches(width_in)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_in * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def _set_paragraph_run_font(run, *, name: str = "Arial", size: int = 10, bold: bool = False, italic: bool = False, color: str = "000000") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def _style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10)

    for style_name, size, bold, color in [
        ("Title", 20, True, "1F3A5F"),
        ("Heading 1", 15, True, "1F3A5F"),
        ("Heading 2", 12, True, "244062"),
        ("Heading 3", 10.5, True, "2F5597"),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)

    if "Code Block" not in [s.name for s in styles]:
        style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Consolas"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        style.font.size = Pt(8.5)


def _add_title_block(doc: Document) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("TranslateIT System Architecture and Engine Reference")
    _set_paragraph_run_font(run, size=20, bold=True, color="1F3A5F")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(10)
    run = subtitle.add_run("Current Experimental baseline with launcher, runtime engines, live pipeline, snapshots, and planning constraints")
    _set_paragraph_run_font(run, size=10, color="4F4F4F")

    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    col_widths = [1.45, 5.95]
    for label, value in [
        ("Project", "TranslateIT"),
        ("Working root", str(PROJECT_ROOT)),
        ("Purpose", "Complete technical reference for AI Research and future planning"),
        ("Status", "Living reference for the current Experimental baseline"),
        ("Audience", "AI Research, engineering planning, implementation follow-up"),
    ]:
        row = table.add_row().cells
        _set_cell_width(row[0], col_widths[0])
        _set_cell_width(row[1], col_widths[1])
        _set_cell_text(row[0], label, bold=True, size=9)
        _set_cell_text(row[1], value, size=9)
        _shade_cell(row[0], "D9E2F3")

    doc.add_paragraph("")


def _is_table_sep(line: str) -> bool:
    return bool(re.fullmatch(r"\s*\|?[\s:\-|\+]+\|?\s*", line)) and "-" in line


def _split_table_row(line: str) -> list[str]:
    stripped = line.strip().strip("|")
    return [cell.strip() for cell in stripped.split("|")]


def _apply_paragraph_format(paragraph, *, before: int = 0, after: int = 3, line_spacing: float = 1.08, keep_together: bool = False) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line_spacing
    paragraph.paragraph_format.keep_together = keep_together


def _add_text_paragraph(doc: Document, text: str, *, style: str = "Normal", bold: bool = False, italic: bool = False, size: float | None = None, color: str = "000000") -> None:
    p = doc.add_paragraph(style=style)
    _apply_paragraph_format(p)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def _add_code_block(doc: Document, lines: list[str]) -> None:
    for line in lines:
        p = doc.add_paragraph(style="Code Block")
        _apply_paragraph_format(p, before=0, after=0, line_spacing=1.0)
        p.paragraph_format.left_indent = Inches(0.2)
        run = p.add_run(line)
        _set_paragraph_run_font(run, name="Consolas", size=8.5, color="222222")


def _add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = len(rows[0])
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.LEFT
    avail_width = 6.9
    col_scores = [max(len(r[i]) for r in rows) for i in range(cols)]
    score_sum = max(sum(col_scores), 1)
    widths = [max(0.8, avail_width * score / score_sum) for score in col_scores]
    for row_values in rows:
        row = table.add_row().cells
        for idx, value in enumerate(row_values):
            cell = row[idx]
            _set_cell_width(cell, widths[idx])
            _set_cell_text(cell, value, bold=(row_values is rows[0]), size=8.5 if cols > 4 else 9)
            if row_values is rows[0]:
                _shade_cell(cell, "D9E2F3")
    doc.add_paragraph("")


def _markdown_to_docx(md_text: str, doc: Document) -> None:
    lines = md_text.splitlines()
    i = 0
    in_code = False
    code_lines: list[str] = []
    pending_paragraph: list[str] = []

    def flush_paragraph() -> None:
        nonlocal pending_paragraph
        if not pending_paragraph:
            return
        text = " ".join(part.strip() for part in pending_paragraph).strip()
        if text:
            _add_text_paragraph(doc, text)
        pending_paragraph = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                _add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                flush_paragraph()
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line.rstrip())
            i += 1
            continue
        if not stripped:
            flush_paragraph()
            i += 1
            continue
        if stripped.startswith("#"):
            flush_paragraph()
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            if level == 1 and text == "TranslateIT System Architecture and Engine Reference":
                i += 1
                continue
            if level == 1:
                _add_text_paragraph(doc, text, style="Heading 1", bold=True, size=15, color="1F3A5F")
            elif level == 2:
                _add_text_paragraph(doc, text, style="Heading 2", bold=True, size=12, color="244062")
            elif level == 3:
                _add_text_paragraph(doc, text, style="Heading 3", bold=True, size=10.5, color="2F5597")
            else:
                _add_text_paragraph(doc, text, bold=True, size=10.5, color="2F5597")
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and _is_table_sep(lines[i + 1]):
            flush_paragraph()
            header = _split_table_row(line)
            table_rows = [header]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_rows.append(_split_table_row(lines[i]))
                i += 1
            _add_markdown_table(doc, table_rows)
            continue
        bullet_match = re.match(r"^-\s+(.*)$", stripped)
        number_match = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if bullet_match:
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet")
            _apply_paragraph_format(p, before=0, after=2, line_spacing=1.05)
            run = p.add_run(bullet_match.group(1))
            _set_paragraph_run_font(run, size=10)
            i += 1
            continue
        if number_match:
            flush_paragraph()
            p = doc.add_paragraph(style="List Number")
            _apply_paragraph_format(p, before=0, after=2, line_spacing=1.05)
            run = p.add_run(number_match.group(2))
            _set_paragraph_run_font(run, size=10)
            i += 1
            continue
        pending_paragraph.append(line)
        i += 1
    flush_paragraph()


def _walk_files(root: Path, predicate=None) -> list[Path]:
    result: list[Path] = []
    for path in root.rglob("*"):
        if not path.is_file():
            continue
        if predicate is not None and not predicate(path):
            continue
        result.append(path)
    return sorted(result)


def _collect_manifest() -> dict[str, object]:
    engine_files = _walk_files(
        PROJECT_ROOT / "EngineData",
        lambda p: "__pycache__" not in p.parts and "ModelData" not in p.parts and p.suffix != ".pyc",
    )
    doc_files = _walk_files(
        PROJECT_ROOT / "DevelopingData" / "DocumentationData" / "SourceDocument",
        lambda p: p.suffix in {".md", ".docx"} and p.name != MANIFEST_PATH.name,
    )
    log_files = _walk_files(PROJECT_ROOT / "UserData" / "LogData", lambda p: p.suffix in {".json", ".txt", ".log", ".md", ".png"})

    snapshots = {}
    for rel in [
        ("audio_settings", PROJECT_ROOT / "UserData" / "CacheData" / "audio_settings.json"),
        ("engine_readiness", PROJECT_ROOT / "UserData" / "LogData" / "engine_readiness_latest.json"),
        ("engine_health", PROJECT_ROOT / "UserData" / "LogData" / "engine_health_latest.json"),
        ("start_stop_lifecycle", PROJECT_ROOT / "UserData" / "LogData" / "start_stop_lifecycle_latest.json"),
        ("version_snapshot", PROJECT_ROOT / "UserData" / "LogData" / "version_snapshot_latest.txt"),
        ("app_crash_latest", PROJECT_ROOT / "UserData" / "LogData" / "app_crash_latest.json"),
    ]:
        key, path = rel
        if path.exists():
            snapshots[key] = {
                "path": str(path.relative_to(PROJECT_ROOT)).replace("\\", "/"),
                "bytes": path.stat().st_size,
                "sha256": _sha256(path),
                "content": json.loads(path.read_text(encoding="utf-8")) if path.suffix == ".json" else path.read_text(encoding="utf-8"),
            }

    source_stats = {
        "engine_data_source": {
            "count": len(engine_files),
            "bytes": sum(p.stat().st_size for p in engine_files),
        },
        "source_document_files": {
            "count": len(doc_files),
            "bytes": sum(p.stat().st_size for p in doc_files),
        },
        "log_files": {
            "count": len(log_files),
            "bytes": sum(p.stat().st_size for p in log_files),
        },
        "audio_segments": {
            "count": len(_walk_files(PROJECT_ROOT / "UserData" / "CacheData" / "audio_segments")),
            "bytes": sum(p.stat().st_size for p in _walk_files(PROJECT_ROOT / "UserData" / "CacheData" / "audio_segments")),
        },
        "session_cache": {
            "count": len(_walk_files(PROJECT_ROOT / "UserData" / "CacheData" / "session_cache")),
            "bytes": sum(p.stat().st_size for p in _walk_files(PROJECT_ROOT / "UserData" / "CacheData" / "session_cache")),
        },
    }
    return {
        "project": {
            "name": "TranslateIT",
            "root": str(PROJECT_ROOT),
            "generated_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        },
        "document": {
            "markdown_path": str(MARKDOWN_PATH.relative_to(PROJECT_ROOT)).replace("\\", "/"),
            "docx_path": str(DOCX_PATH.relative_to(PROJECT_ROOT)).replace("\\", "/"),
            "manifest_path": str(MANIFEST_PATH.relative_to(PROJECT_ROOT)).replace("\\", "/"),
        },
        "runtime_paths": {
            "log_dir": "UserData/LogData",
            "cache_dir": "UserData/CacheData",
            "saved_data_dir": "UserData/SavedData",
            "asr_model_dir": "EngineData/TranscriptEngine/ModelData",
            "translation_model_dir": "EngineData/TranslateEngine/ModelData",
            "launcher_entry_vbs": "TranslateIT.vbs",
            "launcher_bootstrap": "EngineData/LauncherApp/launcher_bootstrap.py",
            "main_window": "EngineData/LauncherApp/app_main.py",
        },
        "app_config_defaults": {
            "app_name": "TranslateIT",
            "language_focus_mode": "ID/EN Focus",
            "use_custom_voice_actor": True,
            "voice_actor_profile_id": "marcel",
            "voice_actor_profiles_root": r"D:\Work\AI Stuff\TranslateIT-ISSUED\DevelopingPack\UserData\SavedData\profiles\default\voices",
            "primary_asr_model": "large-v3-turbo",
            "backup_asr_model": "medium",
            "device": "cuda",
            "compute_type": "float16",
            "source_language": "id",
            "target_language": "en",
            "asr_task": "transcribe",
            "vad_preset": "Headset",
            "translation_engine_name": "local-nllb-distilled",
            "translation_fallback_engine_name": "marianmt-id-en",
            "tts_enabled": False,
            "local_only_mode": True,
            "temperature": 0,
            "beam_size": 1,
            "condition_on_previous_text": False,
            "vad_filter": False,
            "word_timestamps": False,
        },
        "audio_settings_defaults": {
            "input_device_id": None,
            "output_device_id": None,
            "input_sensitivity": "Headset",
            "show_advanced_devices": False,
            "allow_low_but_usable_input": True,
            "auto_play_out_voice": True,
            "use_custom_voice_actor": True,
            "voice_actor_profile_id": "marcel",
            "voice_actor_profiles_root": r"D:\Work\AI Stuff\TranslateIT-ISSUED\DevelopingPack\UserData\SavedData\profiles\default\voices",
        },
        "ui_states": [
            "IDLE",
            "READY",
            "PREPARING",
            "STREAM_CHECK",
            "READY_TO_LISTEN",
            "LISTENING",
            "SPEECH_DETECTED",
            "TRANSCRIBING",
            "TRANSLATING",
            "COMPLETED",
            "PAUSED",
            "STOPPED",
            "ERROR",
        ],
        "capture_modes": [
            "Diagnostic Only",
            "Mock Pipeline",
            "Real ASR + Mock Translation",
            "Real ASR + Real Translation",
        ],
        "launch_chain": {
            "vbs_entry": "TranslateIT.vbs",
            "vbs_behavior": [
                "Create UserData\\LogData if needed",
                "Write launcher_latest.log entries",
                "Default TRANSLATEIT_TTS_BACKEND to sapi_direct_async if unset",
                "Set TRANSLATEIT_LAUNCH_MODE=gui",
                "Set TRANSLATEIT_PROJECT_ROOT",
                "Launch pythonw.exe -m EngineData.LauncherApp.launcher_bootstrap",
            ],
            "bootstrap_behavior": [
                "Change cwd to project root",
                "Inject project root into sys.path",
                "Install global crash recorders",
                "Emit startup snapshots",
                "Run self-test with --self-test",
                "Import and execute EngineData.LauncherApp.app_main",
            ],
            "single_instance_lock": "UserData/CacheData/translateit.lock",
        },
        "source_stats": {
            name: {"count": int(values["count"]), "bytes": int(values["bytes"])} for name, values in source_stats.items()
        },
        "snapshots": snapshots,
        "source_files": [_file_record(path, PROJECT_ROOT) for path in engine_files],
        "documentation_files": [_file_record(path, PROJECT_ROOT) for path in doc_files],
        "log_files": [_file_record(path, PROJECT_ROOT) for path in log_files],
    }


def build_docx() -> None:
    doc = Document()
    _style_document(doc)
    _add_title_block(doc)
    markdown = _read_text(MARKDOWN_PATH)
    lines = markdown.splitlines()
    if lines and lines[0].startswith("# "):
        markdown = "\n".join(lines[1:])
    _markdown_to_docx(markdown, doc)
    doc.save(DOCX_PATH)


def build_manifest() -> None:
    manifest = _collect_manifest()
    MANIFEST_PATH.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")


def main() -> int:
    build_docx()
    build_manifest()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
